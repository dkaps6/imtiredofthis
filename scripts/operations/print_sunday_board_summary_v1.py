#!/usr/bin/env python3
"""Print a plain-text summary of Sunday HAS-EDGE rows from the master workbook.

Read-only reporting only. Does not touch projections, probabilities, EV,
pricing, or the Decision policy in scripts/master_betting_workbook_core_v2.py.
It only filters and prints rows that script already produced, so this is
safe to run against any already-built outputs/NFL_BETTING_MODEL_MASTER.xlsx.
"""
from __future__ import annotations

import argparse
import csv
import sys
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook


def parse_args():
    p = argparse.ArgumentParser()
    p.add_argument("--workbook", default="outputs/NFL_BETTING_MODEL_MASTER.xlsx")
    p.add_argument("--weekday", default="Sunday", help="Python weekday name to filter Kickoff UTC to")
    return p.parse_args()


def main() -> int:
    a = parse_args()
    wb = load_workbook(Path(a.workbook))
    ws = wb["Master Betting Board"]
    headers = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(headers)}

    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        kickoff = r[idx["Kickoff UTC"]]
        signal = r[idx["Snapshot Signal"]]
        if not kickoff or signal != "HAS EDGE":
            continue
        try:
            dt = datetime.fromisoformat(str(kickoff).replace("Z", "+00:00"))
        except ValueError:
            continue
        if dt.astimezone(timezone.utc).strftime("%A") != a.weekday:
            continue
        rows.append({
            "player": r[idx["Player"]],
            "team": r[idx["Team"]],
            "opp": r[idx["Opponent"]],
            "market": r[idx["Market"]],
            "book": r[idx["Book"]],
            "line": r[idx["Vegas Line"]],
            "proj": r[idx["Model Projection"]],
            "side": r[idx["Best Side"]],
            "odds": r[idx["Best Odds"]],
            "prob_edge": r[idx["Probability Edge"]],
            "best_ev": r[idx["Best EV ROI"]],
            "decision": r[idx["Decision"]],
            "kickoff": kickoff,
            "bettable_now": r[idx["Bettable Now"]],
        })

    rows.sort(key=lambda x: (x["best_ev"] if x["best_ev"] is not None else -999), reverse=True)

    print(f"=== {a.weekday} HAS-EDGE rows: {len(rows)} ===")
    writer = csv.writer(sys.stdout)
    writer.writerow([
        "player", "team", "opp", "market", "side", "line", "proj",
        "gap_pct_of_line", "odds", "prob_edge_pct", "best_ev_pct", "kickoff_utc",
    ])
    for r in rows:
        line = r["line"] if isinstance(r["line"], (int, float)) else None
        proj = r["proj"] if isinstance(r["proj"], (int, float)) else None
        gap_pct = f"{abs(proj - line) / line * 100:.1f}" if line and proj else ""
        pe = f"{r['prob_edge']*100:.1f}" if isinstance(r["prob_edge"], (int, float)) else ""
        ev = f"{r['best_ev']*100:.1f}" if isinstance(r["best_ev"], (int, float)) else ""
        writer.writerow([
            r["player"], r["team"], r["opp"], r["market"], r["side"],
            f"{line:.1f}" if line is not None else "",
            f"{proj:.1f}" if proj is not None else "",
            gap_pct, r["odds"], pe, ev, r["kickoff"],
        ])
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
