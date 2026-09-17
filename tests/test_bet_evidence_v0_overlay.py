from pathlib import Path

from openpyxl import Workbook, load_workbook

from scripts.bet_evidence_v0_overlay import apply_overlay, classify_v0


def test_classify_v0_rb_has_no_historical_trust_score():
    authority, evidence, context, note = classify_v0("RB", "Rushing Yards")
    assert authority == "NONE"
    assert evidence == "NO_HISTORICAL_TRUST_SCORE"
    assert context == "NO_RETROSPECTIVE_AUTHORITY"
    assert "retrospective" in note.lower()


def test_classify_v0_wr_receiving_yards_is_descriptive_caution():
    authority, evidence, context, note = classify_v0("WR", "Receiving Yards")
    assert authority == "AUTHORITY_EXACT_AVAILABLE_LIMITED"
    assert evidence == "DESCRIPTIVE_ONLY"
    assert context == "HISTORICALLY_WEAK_DIRECTIONAL_DIAGNOSTIC"
    assert "48.85" in note


def test_overlay_adds_sheet_without_mutating_snapshot_values(tmp_path: Path):
    out = tmp_path / "outputs" / "NFL_BETTING_MODEL_MASTER.xlsx"
    out.parent.mkdir(parents=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Best Snapshot Edges"
    ws.append([
        "Player", "Team", "Opp", "Pos", "Position Source", "Model Role", "Market",
        "Best Book", "Vegas Line", "Model Projection", "Projection-Line", "Best Side",
        "Best Odds", "Model P", "Market P (No-Vig)", "Probability Edge", "Best EV ROI",
        "Snapshot Signal", "Current Availability", "Game Status", "Science Status",
        "Bettable Now", "Decision",
    ])
    ws.append([
        "Example QB", "BUF", "MIA", "QB", "PLAYER_FORM", "QB1", "Passing Yards",
        "Book", 250.5, 266.0, 15.5, "OVER", -110, 0.58, 0.52, 0.06, 0.08,
        "HAS EDGE", "AVAILABLE", "PREGAME", "QUALIFIED", True, "OVER",
    ])
    notes = wb.create_sheet("Lineage & Notes")
    notes.append(["Key", "Value"])
    wb.save(out)

    apply_overlay(str(tmp_path), "outputs/NFL_BETTING_MODEL_MASTER.xlsx")

    rebuilt = load_workbook(out, data_only=False)
    assert "Bet Evidence V0" in rebuilt.sheetnames
    src = rebuilt["Best Snapshot Edges"]
    assert src["I2"].value == 250.5
    assert src["J2"].value == 266.0
    ev = rebuilt["Bet Evidence V0"]
    assert ev["P2"].value == "AUTHORITY_EXACT_AVAILABLE"
    assert ev["Q2"].value == "DESCRIPTIVE_ONLY"
    assert ev["R2"].value == "HEALTHIER_DIRECTIONAL_DIAGNOSTIC"
