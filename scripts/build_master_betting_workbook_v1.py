#!/usr/bin/env python3
"""Publish a downstream-only XLSX master board from the current Full Slate run."""
from __future__ import annotations
import argparse, json, math, os
from pathlib import Path
from collections import defaultdict
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.formatting.rule import FormulaRule, ColorScaleRule
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

NAVY='172554'; SLATE='334155'; WHITE='FFFFFF'; LBLUE='DBEAFE'; GREEN='DCFCE7'; GD='166534'; YELLOW='FEF3C7'; YD='92400E'; RED='FEE2E2'; RD='991B1B'; GRAY='F1F5F9'; GRAYD='475569'
MLABEL={'player_pass_yds':'Passing Yards','player_rush_yds':'Rushing Yards','player_reception_yds':'Receiving Yards','player_receptions':'Receptions','player_rush_reception_yds':'Rush + Rec Yards','player_anytime_td':'Anytime TD'}

def csvp(p):
    try: return pd.read_csv(p,low_memory=False) if p.exists() and p.stat().st_size else pd.DataFrame()
    except Exception: return pd.DataFrame()
def jsonp(p):
    try: return json.load(open(p,encoding='utf-8')) if p.exists() and p.stat().st_size else {}
    except Exception: return {}
def num(x):
    try: x=float(x); return x if math.isfinite(x) else None
    except Exception: return None
def bol(x): return str(x).strip().lower() in {'1','true','yes','y'} or x is True
def imp(o):
    o=num(o)
    if o is None or o==0:return None
    return 100/(o+100) if o>0 else -o/(-o+100)
def ev(p,o):
    p=num(p); o=num(o)
    if p is None or o is None or o==0:return None
    return p*(o/100 if o>0 else 100/abs(o))-(1-p)
def nv(a,b):
    if a is None:return None
    if b is None:return a
    return a/(a+b) if a+b else None

def header(ws):
    for c in ws[1]: c.fill=PatternFill('solid',fgColor=SLATE); c.font=Font(bold=True,color=WHITE); c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
def table(ws,name):
    if ws.max_row<2:return
    t=Table(displayName=name,ref=f'A1:{get_column_letter(ws.max_column)}{ws.max_row}'); t.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True); ws.add_table(t)
def widths(ws,cap=34):
    for cells in ws.columns:
        letter=get_column_letter(cells[0].column); n=max([len(str(c.value)) for c in cells[:80] if c.value is not None] or [8]); ws.column_dimensions[letter].width=min(max(n+2,9),cap)
def frame(wb,name,df,tname):
    ws=wb.create_sheet(name)
    if df.empty: ws.append(['Status','Message']); ws.append(['NO DATA',f'{name} was not available in this run.'])
    else:
        d=df.fillna(''); ws.append(list(map(str,d.columns)))
        for r in d.itertuples(index=False,name=None): ws.append(list(r))
    header(ws); ws.freeze_panes='A2'; widths(ws,38); table(ws,tname)
    return ws

def maps(av,pf,cert,lineage):
    aex={}; ak=defaultdict(list); pex={}; pk=defaultdict(list); tg={}; ln={}
    if not av.empty:
        for r in av.fillna('').to_dict('records'): aex[(str(r.get('player_clean_key','')),str(r.get('team','')))]=r; ak[str(r.get('player_clean_key',''))].append(r)
    if not pf.empty:
        for r in pf.fillna('').to_dict('records'): pex[(str(r.get('player_clean_key','')),str(r.get('team','')))]=r; pk[str(r.get('player_clean_key',''))].append(r)
    if not cert.empty:
        for r in cert.fillna('').to_dict('records'):
            for c in ('away_team','home_team'):
                if r.get(c): tg[str(r[c])]=r
    if not lineage.empty:
        for r in lineage.fillna('').to_dict('records'): ln[(str(r.get('market','')),str(r.get('position_family','')))]=r
    return aex,ak,pex,pk,tg,ln
def find(k,t,ex,keyed):
    if (k,t) in ex:return ex[(k,t)],'EXACT'
    x=keyed.get(k,[]); return (x[0],'UNIQUE_PLAYER_KEY') if len(x)==1 else (None,'UNMATCHED')
def fam(pos,m):
    if m=='player_anytime_td':return 'ALL'
    p=str(pos).upper().strip().replace('HB','RB').replace('TB','RB')
    if p in {'RB','FB'}:return 'RB/FB'
    if p in {'QB','WR','TE'}:return p
    return 'QB' if m=='player_pass_yds' else ('RB/FB' if m=='player_rush_reception_yds' else p or 'UNKNOWN')

def pair(pr,av,pf,cert,lineage):
    if pr.empty:return []
    for c in ['player_clean_key','book_title','vegas_line','model_proj','model_sd','vegas_over_odds','vegas_under_odds','fair_prob']:
        if c not in pr.columns: pr[c]=''
    aex,ak,pex,pk,tg,ln=maps(av,pf,cert,lineage); out=[]
    gc=['event_id','player_clean_key','team','opponent','source_market','book','vegas_line']
    for _,g in pr.groupby(gc,dropna=False,sort=False):
        b=g.iloc[0].to_dict(); ov=g[g.side.astype(str).str.upper().eq('OVER')]; un=g[g.side.astype(str).str.upper().eq('UNDER')]; O=ov.iloc[0].to_dict() if len(ov) else {}; U=un.iloc[0].to_dict() if len(un) else {}
        k=str(b.get('player_clean_key','')); t=str(b.get('team','')); a,am=find(k,t,aex,ak); p,_=find(k,t,pex,pk); a=a or {}; p=p or {}; gm=tg.get(t,{})
        pos=str(p.get('position_group') or p.get('position') or a.get('position_group') or a.get('position') or 'UNKNOWN'); mr=str(p.get('model_role') or p.get('role') or ''); dr=str(a.get('depth_chart_role') or a.get('raw_depth_role') or '')
        avail='UNAVAILABLE' if bol(a.get('definitive_unavailable')) else str(a.get('final_availability_state') or ('UNMATCHED_CURRENT_ROSTER' if not a else 'AVAILABLE'))
        m=str(b.get('source_market','')); li=ln.get((m,fam(pos,m)),ln.get((m,'ALL'),{})); sci=str(li.get('scientific_status','LINEAGE_NOT_RESOLVED'))
        op=num(O.get('fair_prob')); up=num(U.get('fair_prob')); oo=num(b.get('vegas_over_odds')); uo=num(b.get('vegas_under_odds')); oe=ev(op,oo); ue=ev(up,uo); side='' if oe is None and ue is None else ('OVER' if ue is None or (oe is not None and oe>=ue) else 'UNDER')
        ro,ru=imp(oo),imp(uo); no,nu=nv(ro,ru),nv(ru,ro); bmp=op if side=='OVER' else up if side=='UNDER' else None; bmk=no if side=='OVER' else nu if side=='UNDER' else None
        out.append(dict(event_id=str(b.get('event_id','')),game_id=str(gm.get('game_id','')),kickoff=str(gm.get('kickoff_utc','')),player=str(b.get('player','')),key=k,team=t,opp=str(b.get('opponent','')),pos=pos,model_role=mr,depth_role=dr,availability=avail,authority=str(a.get('availability_authority','')),match=am,game_status=str(gm.get('certification_state','UNMATCHED_GAME')),game_eligible=bol(gm.get('production_eligible')),market=m,market_label=MLABEL.get(m,m),book=str(b.get('book_title') or b.get('book','')),line=num(b.get('vegas_line')),proj=num(b.get('model_proj')),sd=num(b.get('model_sd')),over_odds=oo,under_odds=uo,p_over=op,p_under=up,raw_over=ro,raw_under=ru,nv_over=no,nv_under=nu,ev_over=oe,ev_under=ue,best_side=side,best_ev=oe if side=='OVER' else ue if side=='UNDER' else None,best_model_p=bmp,best_market_p=bmk,prob_edge=None if bmp is None or bmk is None else bmp-bmk,science=sci,mean_owner=str(li.get('final_mean_owner','')),limitation=str(li.get('known_limitation',''))))
    return out

def signal(r,current):
    if 'NOT_DEDICATED' in r['science'] or r['science']=='LINEAGE_NOT_RESOLVED':return 'RESEARCH ONLY','BLOCKED'
    if r['game_status']=='KICKED_OFF_LOCKED' or r['availability'] in {'UNAVAILABLE','UNMATCHED_CURRENT_ROSTER'}:return 'BLOCKED','BLOCKED'
    e=r['best_ev']; q=r['prob_edge']; s='STRONG EDGE' if e is not None and q is not None and e>=.05 and q>=.03 else ('LEAN EDGE' if e is not None and e>0 else 'NO EDGE')
    if not current:return s,'NO LIVE ODDS'
    return s,('PLAY '+r['best_side'] if s=='STRONG EDGE' else ('LEAN '+r['best_side'] if s=='LEAN EDGE' else 'PASS'))

def build(a):
    root=Path(a.root); out=root/a.out; out.parent.mkdir(parents=True,exist_ok=True)
    pr=csvp(root/'outputs/props_priced_clean.csv'); av=csvp(root/'data/current_player_availability.csv'); cert=csvp(root/'data/current_player_availability_game_certification.csv'); pf=csvp(root/'data/player_form.csv'); lin=csvp(root/'data/market_model_lineage_current.csv'); live=jsonp(root/'data/live_odds_status.json')
    requested=bol(a.fetch_live_odds); status='CURRENT' if requested and bol(live.get('available')) and not pr.empty else ('NO_LIVE_ODDS_REQUESTED' if not requested else 'NO_ACTIVE_MARKETS'); rows=pair(pr,av,pf,cert,lin); current=status=='CURRENT'
    wb=Workbook(); wb.calculation.fullCalcOnLoad=True; wb.calculation.forceFullCalc=True; wb.calculation.calcMode='auto'; ws=wb.active; ws.title='Dashboard'; ws.merge_cells('A1:H2'); ws['A1']='NFL BETTING MODEL MASTER'; ws['A1'].fill=PatternFill('solid',fgColor=NAVY); ws['A1'].font=Font(bold=True,color=WHITE,size=16)
    ws.merge_cells('A4:H5'); ws['A4']=('CURRENT sportsbook prices are attached to this same Full Slate run.' if current else ('Live odds were not requested; football outputs are current and betting decisions are suppressed.' if not requested else 'Live odds were requested but no active pricing was available; betting decisions are suppressed.')); ws['A4'].fill=PatternFill('solid',fgColor=GREEN if current else YELLOW); ws['A4'].font=Font(bold=True,color=GD if current else YD); ws['A4'].alignment=Alignment(wrap_text=True)
    kpis=[('Sportsbook Offers',len(rows)),('Unique Player-Markets',len({(r['key'],r['market']) for r in rows})),('Priced Players',len({r['key'] for r in rows if r['key']})),('Eligible Games',sum(bol(x) for x in cert.get('production_eligible',pd.Series(dtype=object)))),('Locked Games',sum(str(x)=='KICKED_OFF_LOCKED' for x in cert.get('certification_state',pd.Series(dtype=object)))),('Availability Rows',len(av))]
    for i,(lab,val) in enumerate(kpis): rr=7+(i//3)*3; cc=1+(i%3)*2; ws.cell(rr,cc,lab).fill=PatternFill('solid',fgColor=LBLUE); ws.cell(rr,cc).font=Font(bold=True,color=NAVY); ws.cell(rr+1,cc,val).font=Font(bold=True,color=NAVY,size=16)
    ws['A14']='Pricing Status'; ws['B14']=status; ws.merge_cells('A18:H21'); ws['A18']='Sportsbook is downstream only. This workbook packages the current run; it does not feed lines or probabilities back into football projections. PLAY/LEAN labels are display gates, not a calibrated staking policy.'; ws['A18'].fill=PatternFill('solid',fgColor=GRAY); ws['A18'].font=Font(color=GRAYD); ws['A18'].alignment=Alignment(wrap_text=True)
    for c in range(1,9): ws.column_dimensions[get_column_letter(c)].width=20

    me=wb.create_sheet('Master Betting Board'); H=['Event ID','Game ID','Kickoff UTC','Player','Player Key','Team','Opponent','Position','Model Role','Depth Role','Current Availability','Availability Authority','Roster Match','Game Status','Game Eligible','Market','Book','Vegas Line','Model Projection','Model SD','Projection - Line','Over Odds','Under Odds','Model P Over','Model P Under','No-Vig P Over','No-Vig P Under','EV ROI Over','EV ROI Under','Best Side','Best Odds','Best Model P','Best Market P','Probability Edge','Best EV ROI','Snapshot Signal','Science Status','Bettable Now','Decision','Mean Owner / Limitation']; me.append(H)
    for r in rows:
        sg,dec=signal(r,current); me.append([r['event_id'],r['game_id'],r['kickoff'],r['player'],r['key'],r['team'],r['opp'],r['pos'],r['model_role'],r['depth_role'],r['availability'],r['authority'],r['match'],r['game_status'],r['game_eligible'],r['market_label'],r['book'],r['line'],r['proj'],r['sd'],None,r['over_odds'],r['under_odds'],r['p_over'],r['p_under'],r['nv_over'],r['nv_under'],r['ev_over'],r['ev_under'],r['best_side'],r['over_odds'] if r['best_side']=='OVER' else r['under_odds'] if r['best_side']=='UNDER' else '',r['best_model_p'],r['best_market_p'],r['prob_edge'],r['best_ev'],sg,r['science'],current and r['game_eligible'] and sg not in {'BLOCKED','RESEARCH ONLY'},dec,(r['mean_owner']+(' | '+r['limitation'] if r['limitation'] else ''))[:32000]]); rr=me.max_row; me.cell(rr,21,f'=IF(OR(S{rr}="",R{rr}=""),"",S{rr}-R{rr})')
    header(me); me.freeze_panes='D2'; widths(me,28); me.column_dimensions['AN'].width=70; table(me,'MasterBettingBoard')
    for c in ['X','Y','Z','AA','AB','AC','AF','AG','AH','AI']:
        for cell in me[c][1:]: cell.number_format='0.0%'
    if me.max_row>1:
        me.conditional_formatting.add(f'AJ2:AJ{me.max_row}',FormulaRule(formula=['$AJ2="STRONG EDGE"'],fill=PatternFill('solid',fgColor=GREEN),font=Font(color=GD,bold=True))); me.conditional_formatting.add(f'AJ2:AJ{me.max_row}',FormulaRule(formula=['$AJ2="LEAN EDGE"'],fill=PatternFill('solid',fgColor=YELLOW),font=Font(color=YD,bold=True))); me.conditional_formatting.add(f'AJ2:AJ{me.max_row}',FormulaRule(formula=['OR($AJ2="BLOCKED",$AJ2="RESEARCH ONLY")'],fill=PatternFill('solid',fgColor=RED),font=Font(color=RD,bold=True))); me.conditional_formatting.add(f'AI2:AI{me.max_row}',ColorScaleRule(start_type='min',start_color='FECACA',mid_type='percentile',mid_value=50,mid_color='FEF3C7',end_type='max',end_color='BBF7D0'))

    be=wb.create_sheet('Best Snapshot Edges'); BH=['Player','Team','Opp','Pos','Model Role','Market','Best Book','Vegas Line','Model Projection','Projection-Line','Best Side','Best Odds','Model P','Market P (No-Vig)','Probability Edge','Best EV ROI','Snapshot Signal','Current Availability','Game Status','Science Status','Bettable Now','Decision']; be.append(BH); best={}
    for r in rows:
        k=(r['key'],r['market']); sc=r['best_ev'] if r['best_ev'] is not None else -999
        if k not in best or sc>(best[k]['best_ev'] if best[k]['best_ev'] is not None else -999):best[k]=r
    def rk(r): return (('NOT_DEDICATED' in r['science'] or r['science']=='LINEAGE_NOT_RESOLVED'),r['game_status']=='KICKED_OFF_LOCKED' or r['availability'] in {'UNAVAILABLE','UNMATCHED_CURRENT_ROSTER'},-(r['best_ev'] if r['best_ev'] is not None else -999))
    for r in sorted(best.values(),key=rk): sg,dec=signal(r,current); odds=r['over_odds'] if r['best_side']=='OVER' else r['under_odds'] if r['best_side']=='UNDER' else ''; be.append([r['player'],r['team'],r['opp'],r['pos'],r['model_role'],r['market_label'],r['book'],r['line'],r['proj'],None,r['best_side'],odds,r['best_model_p'],r['best_market_p'],r['prob_edge'],r['best_ev'],sg,r['availability'],r['game_status'],r['science'],current and r['game_eligible'] and sg not in {'BLOCKED','RESEARCH ONLY'},dec]); rr=be.max_row; be.cell(rr,10,f'=IF(OR(I{rr}="",H{rr}=""),"",I{rr}-H{rr})')
    header(be); be.freeze_panes='B2'; widths(be,28); table(be,'BestSnapshotEdges')
    for c in ['M','N','O','P']:
        for cell in be[c][1:]:cell.number_format='0.0%'
    if be.max_row>1:
        be.conditional_formatting.add(f'Q2:Q{be.max_row}',FormulaRule(formula=['$Q2="STRONG EDGE"'],fill=PatternFill('solid',fgColor=GREEN),font=Font(color=GD,bold=True))); be.conditional_formatting.add(f'Q2:Q{be.max_row}',FormulaRule(formula=['$Q2="LEAN EDGE"'],fill=PatternFill('solid',fgColor=YELLOW),font=Font(color=YD,bold=True)))

    frame(wb,'Game Certification',cert,'GameCertification'); avv=av.copy()
    if not avv.empty and not pf.empty and {'player_clean_key','team'}.issubset(avv.columns) and {'player_clean_key','team'}.issubset(pf.columns):
        rc=next((c for c in ('model_role','role') if c in pf.columns),None)
        if rc: avv=avv.merge(pf[['player_clean_key','team',rc]].drop_duplicates(['player_clean_key','team']).rename(columns={rc:'current_model_role'}),on=['player_clean_key','team'],how='left')
    frame(wb,'Availability & Roles',avv,'AvailabilityRoles'); frame(wb,'Market Science',lin,'MarketScience')
    ln=wb.create_sheet('Lineage & Notes'); notes=[('Workbook builder','scripts/build_master_betting_workbook_v1.py'),('GitHub Run ID',a.run_id),('GitHub SHA',a.sha),('GitHub Ref',a.ref_name),('Live odds requested',requested),('Pricing status',status),('Live odds status',live.get('status','')),('Odds API refetched',live.get('odds_api_refetched','')),('Architecture','Sportsbook is downstream only; workbook generation cannot alter model projections.'),('No-odds behavior','Workbook still publishes, but betting decisions are suppressed.')]; ln.append(['NFL BETTING MODEL MASTER — RUN LINEAGE','']); ln.merge_cells('A1:B1'); ln['A1'].fill=PatternFill('solid',fgColor=NAVY); ln['A1'].font=Font(bold=True,color=WHITE,size=15)
    for x,y in notes: ln.append([x,y]); ln.cell(ln.max_row,1).fill=PatternFill('solid',fgColor=LBLUE); ln.cell(ln.max_row,1).font=Font(bold=True,color=NAVY)
    ln.column_dimensions['A'].width=34; ln.column_dimensions['B'].width=90
    wb._sheets=[wb[n] for n in ['Dashboard','Best Snapshot Edges','Master Betting Board','Game Certification','Availability & Roles','Market Science','Lineage & Notes']]; wb.save(out)
    print(json.dumps({'disposition':'MASTER_BETTING_WORKBOOK_PUBLISHED','output':str(out),'pricing_status':status,'priced_offer_rows':len(rows),'player_market_rows':len({(r['key'],r['market']) for r in rows}),'availability_rows':len(av),'game_certification_rows':len(cert)},sort_keys=True))

def args():
    p=argparse.ArgumentParser(); p.add_argument('--root',default='.'); p.add_argument('--out',default='outputs/NFL_BETTING_MODEL_MASTER.xlsx'); p.add_argument('--run-id',default=os.getenv('GITHUB_RUN_ID','')); p.add_argument('--sha',default=os.getenv('GITHUB_SHA','')); p.add_argument('--ref-name',default=os.getenv('GITHUB_REF_NAME','')); p.add_argument('--fetch-live-odds',default=os.getenv('FETCH_LIVE_ODDS','false')); return p.parse_args()
if __name__=='__main__':build(args())
