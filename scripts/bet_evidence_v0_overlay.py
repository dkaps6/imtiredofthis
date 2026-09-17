#!/usr/bin/env python3
"""Add a conservative historical-evidence V0 sheet to the Full Slate workbook.

This module is deliberately downstream-only. It never changes football projections,
probabilities, sportsbook prices, EV, signal, eligibility, or model routing. It only
adds transparent historical context to the already-built workbook.

V0 is NOT the conditional historical-analog engine. It exposes only currently
verified authority coverage and coarse authority-exact directional diagnostics.
Nothing in this sheet is a staking instruction or a validated bet-selection rule.
"""
from __future__ import annotations

import argparse
import json
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

NAVY = "172554"
SLATE = "334155"
WHITE = "FFFFFF"
LBLUE = "DBEAFE"
YELLOW = "FEF3C7"
YD = "92400E"
RED = "FEE2E2"
RD = "991B1B"
GRAY = "F1F5F9"
GRAYD = "475569"

SHEET_NAME = "Bet Evidence V0"


def classify_v0(position: str, market: str) -> tuple[str, str, str, str]:
    """Return authority status, evidence class, context label, and scoped note.

    These are intentionally coarse, pre-existing diagnostics only. They must not be
    interpreted as player/matchup analog support.
    """
    pos = str(position or "").upper().strip()
    mkt = str(market or "").strip()

    if pos in {"RB", "FB"} and mkt in {"Rushing Yards", "Rush + Rec Yards"}:
        return (
            "NONE",
            "NO_HISTORICAL_TRUST_SCORE",
            "NO_RETROSPECTIVE_AUTHORITY",
            "Current RB rushing authority has no legitimate authority-exact retrospective Vegas grade. Do not inherit the old base ensemble track record.",
        )

    if pos == "QB" and mkt == "Passing Yards":
        return (
            "AUTHORITY_EXACT_AVAILABLE",
            "DESCRIPTIVE_ONLY",
            "HEALTHIER_DIRECTIONAL_DIAGNOSTIC",
            "2024 authority-exact QB1 passing-yards directional result was 54.21%. This is descriptive context only; no validated nightly selection rule or matchup-analog rule is approved yet.",
        )

    if pos == "WR" and mkt == "Receiving Yards":
        return (
            "AUTHORITY_EXACT_AVAILABLE_LIMITED",
            "DESCRIPTIVE_ONLY",
            "HISTORICALLY_WEAK_DIRECTIONAL_DIAGNOSTIC",
            "2024 authority-exact WR receiving-yards direction was weak: WR1 48.85% and matched WR2+ about 48.58%. Treat as a caution diagnostic, not an automatic fade rule.",
        )

    if pos == "WR" and mkt == "Receptions":
        return (
            "AUTHORITY_EXACT_AVAILABLE_LIMITED",
            "DESCRIPTIVE_ONLY",
            "HEALTHIER_DIRECTIONAL_DIAGNOSTIC",
            "2024 authority-exact WR1 receptions directional result was 54.33%. This is descriptive context only; WR-R15 historical scope is limited and no validated nightly selection rule is approved yet.",
        )

    if pos == "TE" and mkt == "Receiving Yards":
        return (
            "AUTHORITY_EXACT_AVAILABLE",
            "DESCRIPTIVE_ONLY",
            "HEALTHIER_DIRECTIONAL_DIAGNOSTIC",
            "Authority-exact TE1 receiving-yards directional result was 54.70% in the cited benchmark. Descriptive only; not a validated nightly selection rule.",
        )

    if pos == "TE" and mkt == "Receptions":
        return (
            "AUTHORITY_EXACT_AVAILABLE",
            "DESCRIPTIVE_ONLY",
            "HEALTHIER_DIRECTIONAL_DIAGNOSTIC",
            "Authority-exact TE1 receptions directional result was 55.73% in the cited benchmark. Descriptive only; not a validated nightly selection rule.",
        )

    return (
        "NO_APPROVED_V0_MAPPING",
        "NO_EVIDENCE",
        "NO_APPROVED_HISTORICAL_MAPPING",
        "No approved V0 historical context mapping exists for this position/market. This is not evidence against the bet; it means V0 should abstain.",
    )


def _header(ws) -> None:
    for c in ws[1]:
        c.fill = PatternFill("solid", fgColor=SLATE)
        c.font = Font(bold=True, color=WHITE)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _widths(ws, cap: int = 42) -> None:
    for cells in ws.columns:
        letter = get_column_letter(cells[0].column)
        n = max([len(str(c.value)) for c in cells[:80] if c.value is not None] or [8])
        ws.column_dimensions[letter].width = min(max(n + 2, 9), cap)


def parse_args():
    p = argparse.ArgumentParser(add_help=False)
    p.add_argument("--root", default=".")
    p.add_argument("--out", default="outputs/NFL_BETTING_MODEL_MASTER.xlsx")
    args, _ = p.parse_known_args()
    return args


def apply_overlay(root: str = ".", out: str = "outputs/NFL_BETTING_MODEL_MASTER.xlsx") -> dict:
    root_path = Path(root)
    workbook_path = root_path / out
    if not workbook_path.exists() or not workbook_path.stat().st_size:
        raise RuntimeError(f"bet evidence V0: workbook missing: {workbook_path}")

    wb = load_workbook(workbook_path)
    if "Best Snapshot Edges" not in wb.sheetnames:
        raise RuntimeError("bet evidence V0: Best Snapshot Edges sheet missing")

    if SHEET_NAME in wb.sheetnames:
        del wb[SHEET_NAME]

    src = wb["Best Snapshot Edges"]
    headers = [str(c.value or "") for c in src[1]]
    index = {name: i for i, name in enumerate(headers)}
    required = {"Player", "Team", "Opp", "Pos", "Market", "Best Book", "Vegas Line", "Model Projection", "Best Side", "Best Odds", "Probability Edge", "Best EV ROI", "Snapshot Signal", "Bettable Now", "Decision"}
    missing = sorted(required - set(index))
    if missing:
        raise RuntimeError(f"bet evidence V0: Best Snapshot Edges missing columns: {missing}")

    ws = wb.create_sheet(SHEET_NAME, 2)
    out_headers = [
        "Player", "Team", "Opp", "Pos", "Market", "Best Book", "Vegas Line",
        "Model Projection", "Best Side", "Best Odds", "Probability Edge", "Best EV ROI",
        "Snapshot Signal", "Bettable Now", "Decision", "Historical Authority Status",
        "Evidence Class", "V0 Historical Context", "Historical Note",
    ]
    ws.append(out_headers)

    counts = {}
    row_count = 0
    for values in src.iter_rows(min_row=2, values_only=True):
        row = {name: values[i] if i < len(values) else None for name, i in index.items()}
        authority, evidence, context, note = classify_v0(row.get("Pos"), row.get("Market"))
        ws.append([
            row.get("Player"), row.get("Team"), row.get("Opp"), row.get("Pos"), row.get("Market"),
            row.get("Best Book"), row.get("Vegas Line"), row.get("Model Projection"), row.get("Best Side"),
            row.get("Best Odds"), row.get("Probability Edge"), row.get("Best EV ROI"), row.get("Snapshot Signal"),
            row.get("Bettable Now"), row.get("Decision"), authority, evidence, context, note,
        ])
        row_count += 1
        counts[evidence] = counts.get(evidence, 0) + 1

    _header(ws)
    ws.freeze_panes = "B2"
    _widths(ws)
    ws.column_dimensions["S"].width = 85
    for col in ["K", "L"]:
        for cell in ws[col][1:]:
            cell.number_format = "0.0%"

    if ws.max_row > 1:
        tab = Table(displayName="BetEvidenceV0", ref=f"A1:S{ws.max_row}")
        tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(tab)
        for r in range(2, ws.max_row + 1):
            evidence = str(ws.cell(r, 17).value or "")
            context = str(ws.cell(r, 18).value or "")
            if evidence == "NO_HISTORICAL_TRUST_SCORE":
                for c in range(16, 20):
                    ws.cell(r, c).fill = PatternFill("solid", fgColor=RED)
                    ws.cell(r, c).font = Font(color=RD, bold=(c == 17))
            elif context == "HISTORICALLY_WEAK_DIRECTIONAL_DIAGNOSTIC":
                for c in range(16, 20):
                    ws.cell(r, c).fill = PatternFill("solid", fgColor=YELLOW)
                    ws.cell(r, c).font = Font(color=YD, bold=(c in {17, 18}))
            elif evidence == "NO_EVIDENCE":
                for c in range(16, 20):
                    ws.cell(r, c).fill = PatternFill("solid", fgColor=GRAY)
                    ws.cell(r, c).font = Font(color=GRAYD, bold=(c == 17))
            else:
                ws.cell(r, 17).fill = PatternFill("solid", fgColor=LBLUE)
                ws.cell(r, 17).font = Font(color=NAVY, bold=True)

    if "Lineage & Notes" in wb.sheetnames:
        notes = wb["Lineage & Notes"]
        notes.append(["Bet Evidence V0", "Downstream descriptive overlay only. It does not change football projections, probabilities, EV, market prices, eligibility, Snapshot Signal, or Decision."])
        notes.cell(notes.max_row, 1).fill = PatternFill("solid", fgColor=LBLUE)
        notes.cell(notes.max_row, 1).font = Font(bold=True, color=NAVY)
        notes.append(["Bet Evidence V0 limitation", "This is NOT the conditional historical-analog engine. It contains only scoped authority-exact directional diagnostics already documented in Issue #535. DESCRIPTIVE_ONLY is not a bet recommendation; NO_HISTORICAL_TRUST_SCORE means the current authority lacks a legitimate retrospective Vegas grade."])
        notes.cell(notes.max_row, 1).fill = PatternFill("solid", fgColor=LBLUE)
        notes.cell(notes.max_row, 1).font = Font(bold=True, color=NAVY)

    wb.save(workbook_path)

    audit = {
        "disposition": "BET_EVIDENCE_V0_OVERLAY_ADDED",
        "workbook": str(workbook_path),
        "rows": row_count,
        "evidence_counts": counts,
        "downstream_only": True,
        "changes_projection": False,
        "changes_probability": False,
        "changes_ev": False,
        "changes_signal": False,
        "conditional_analog_engine": False,
    }
    audit_path = root_path / "data" / "bet_evidence_v0_audit.json"
    audit_path.parent.mkdir(parents=True, exist_ok=True)
    audit_path.write_text(json.dumps(audit, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    print(json.dumps(audit, sort_keys=True))
    return audit


def main() -> int:
    a = parse_args()
    apply_overlay(a.root, a.out)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
