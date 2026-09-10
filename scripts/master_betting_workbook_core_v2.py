#!/usr/bin/env python3
"""Core publisher for the downstream-only Full Slate master betting workbook."""
from __future__ import annotations

import argparse
import json
import math
import os
import re
from collections import defaultdict
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule, FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

NAVY='172554'; SLATE='334155'; WHITE='FFFFFF'; LBLUE='DBEAFE'
GREEN='DCFCE7'; GD='166534'; YELLOW='FEF3C7'; YD='92400E'
RED='FEE2E2'; RD='991B1B'; GRAY='F1F5F9'; GRAYD='475569'
MLABEL={
    'player_pass_yds':'Passing Yards',
    'player_rush_yds':'Rushing Yards',
    'player_reception_yds':'Receiving Yards',
    'player_receptions':'Receptions',
    'player_rush_reception_yds':'Rush + Rec Yards',
    'player_anytime_td':'Anytime TD',
}


def csvp(path: Path) -> pd.DataFrame:
    try:
        return pd.read_csv(path, low_memory=False) if path.exists() and path.stat().st_size else pd.DataFrame()
    except Exception:
        return pd.DataFrame()


def jsonp(path: Path) -> dict:
    try:
        return json.load(open(path, encoding='utf-8')) if path.exists() and path.stat().st_size else {}
    except Exception:
        return {}


def num(value):
    try:
        value=float(value)
        return value if math.isfinite(value) else None
    except Exception:
        return None


def bol(value) -> bool:
    return value is True or str(value).strip().lower() in {'1','true','yes','y'}


def implied_prob(odds):
    odds=num(odds)
    if odds is None or odds == 0:
        return None
    return 100/(odds+100) if odds > 0 else -odds/(-odds+100)


def ev_roi(prob, odds):
    prob=num(prob); odds=num(odds)
    if prob is None or odds is None or odds == 0:
        return None
    return prob*(odds/100 if odds > 0 else 100/abs(odds))-(1-prob)


def no_vig(a,b):
    if a is None:
        return None
    if b is None:
        return a
    return a/(a+b) if a+b else None


def identity_alias(key: str) -> str:
    """Suffix-insensitive player identity used only for workbook joins.

    The production football identity remains untouched. This handles sportsbook
    names such as Marvin Harrison Jr. against Ourlads/player-form keys such as
    marvinharrison.
    """
    s=re.sub(r'[^a-z0-9]','',str(key).lower())
    for suffix in ('iii','ii','iv','jr','sr','v'):
        if s.endswith(suffix) and len(s) > len(suffix)+2:
            return s[:-len(suffix)]
    return s


def normalize_position(value: str) -> str:
    p=str(value or '').upper().strip()
    if p in {'HB','TB'}:
        return 'RB'
    if p in {'LWR','RWR','SWR'}:
        return 'WR'
    if p in {'QB','RB','FB','WR','TE'}:
        return p
    return ''


def _add_identity_maps(df, exact, alias_team, alias_global):
    if df.empty or 'player_clean_key' not in df.columns:
        return
    for row in df.fillna('').to_dict('records'):
        key=str(row.get('player_clean_key','')); team=str(row.get('team',''))
        if not key:
            continue
        exact[(key,team)].append(row)
        a=identity_alias(key)
        alias_team[(a,team)].append(row)
        alias_global[a].append(row)


def _unique(rows):
    return rows[0] if len(rows)==1 else None


def build_identity_sources(av, pf, roles, season_totals, game_logs):
    src={}
    for name,df in [('player_form',pf),('availability',av),('roles',roles)]:
        exact=defaultdict(list); at=defaultdict(list); ag=defaultdict(list)
        _add_identity_maps(df,exact,at,ag)
        src[name]=(exact,at,ag)

    historical=defaultdict(set)
    if not season_totals.empty:
        for r in season_totals.fillna('').to_dict('records'):
            key=str(r.get('player_clean_key','') or r.get('identity_base_name_key',''))
            p=normalize_position(r.get('historical_position',''))
            if key and p:
                historical[identity_alias(key)].add(p)
    if not game_logs.empty:
        for r in game_logs.fillna('').to_dict('records'):
            key=str(r.get('player_clean_key','') or r.get('identity_base_name_key',''))
            p=normalize_position(r.get('position_group','') or r.get('position',''))
            if key and p:
                historical[identity_alias(key)].add(p)
    return src,historical


def lookup_identity(name, key, team, src):
    """Return current row + provenance without changing production identity."""
    exact,at,ag=src[name]
    r=_unique(exact.get((key,team),[]))
    if r:
        return r,'EXACT'
    a=identity_alias(key)
    r=_unique(at.get((a,team),[]))
    if r:
        return r,'SUFFIX_ALIAS_TEAM'
    r=_unique(ag.get(a,[]))
    if r:
        return r,'SUFFIX_ALIAS_UNIQUE'
    return None,'UNMATCHED'


def infer_position_from_model(row) -> tuple[str,str]:
    role=str(row.get('rules_role','') or '').upper().strip()
    for prefix,pos in [('QB','QB'),('RB','RB'),('FB','FB'),('WR','WR'),('TE','TE')]:
        if role.startswith(prefix):
            return pos,'MODEL_RULE_ROLE'
    if bol(row.get('qb_synthesis_applied')):
        return 'QB','QB_SYNTHESIS_ROUTE'
    if bol(row.get('rb_synthesis_applied')):
        return 'RB','RB_SYNTHESIS_ROUTE'
    market=str(row.get('source_market',''))
    if market == 'player_pass_yds':
        return 'QB','MARKET_DETERMINISTIC'
    return '',''


def resolve_position(base, avrow, pfrow, rolesrow, historical):
    for source,row in [('PLAYER_FORM',pfrow),('AVAILABILITY',avrow),('OURLADS_ROLE',rolesrow)]:
        if row:
            p=normalize_position(row.get('position_group','') or row.get('position','') or row.get('historical_position',''))
            if p:
                return p,source
    hist=historical.get(identity_alias(base.get('player_clean_key','')),set())
    if len(hist)==1:
        return next(iter(hist)),'HISTORICAL_PLAYER_POSITION'
    p,source=infer_position_from_model(base)
    if p:
        return p,source
    return 'UNRESOLVED','UNRESOLVED_POSITION'


def market_family(position, market):
    if market=='player_anytime_td':
        return 'ALL'
    p=normalize_position(position)
    if p in {'RB','FB'}:
        return 'RB/FB'
    if p in {'QB','WR','TE'}:
        return p
    if market=='player_pass_yds':
        return 'QB'
    return 'UNKNOWN'


def header(ws):
    for c in ws[1]:
        c.fill=PatternFill('solid',fgColor=SLATE)
        c.font=Font(bold=True,color=WHITE)
        c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)


def add_table(ws,name):
    if ws.max_row < 2:
        return
    t=Table(displayName=name,ref=f'A1:{get_column_letter(ws.max_column)}{ws.max_row}')
    t.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True)
    ws.add_table(t)


def widths(ws,cap=34):
    for cells in ws.columns:
        letter=get_column_letter(cells[0].column)
        n=max([len(str(c.value)) for c in cells[:80] if c.value is not None] or [8])
        ws.column_dimensions[letter].width=min(max(n+2,9),cap)


def frame(wb,name,df,tname):
    ws=wb.create_sheet(name)
    if df.empty:
        ws.append(['Status','Message']); ws.append(['NO DATA',f'{name} was not available in this run.'])
    else:
        d=df.fillna(''); ws.append(list(map(str,d.columns)))
        for row in d.itertuples(index=False,name=None):
            ws.append(list(row))
    header(ws); ws.freeze_panes='A2'; widths(ws,38); add_table(ws,tname)
    return ws


def build_offer_rows(pr,av,pf,roles,cert,lineage,season_totals,game_logs):
    if pr.empty:
        return []
    for c in ['player_clean_key','book_title','vegas_line','model_proj','model_sd','vegas_over_odds','vegas_under_odds','fair_prob','rules_role','qb_synthesis_applied','rb_synthesis_applied']:
        if c not in pr.columns:
            pr[c]=''
    src,historical=build_identity_sources(av,pf,roles,season_totals,game_logs)
    team_game={}
    if not cert.empty:
        for r in cert.fillna('').to_dict('records'):
            for c in ('away_team','home_team'):
                if r.get(c): team_game[str(r[c])]=r
    lin={}
    if not lineage.empty:
        for r in lineage.fillna('').to_dict('records'):
            lin[(str(r.get('market','')),str(r.get('position_family','')))]=r

    rows=[]
    group_cols=['event_id','player_clean_key','team','opponent','source_market','book','vegas_line']
    for _,g in pr.groupby(group_cols,dropna=False,sort=False):
        b=g.iloc[0].to_dict()
        overs=g[g.side.astype(str).str.upper().eq('OVER')]
        unders=g[g.side.astype(str).str.upper().eq('UNDER')]
        O=overs.iloc[0].to_dict() if len(overs) else {}
        U=unders.iloc[0].to_dict() if len(unders) else {}
        key=str(b.get('player_clean_key','')); team=str(b.get('team',''))
        avrow,avmatch=lookup_identity('availability',key,team,src)
        pfrow,pfmatch=lookup_identity('player_form',key,team,src)
        rolesrow,rolesmatch=lookup_identity('roles',key,team,src)
        gm=team_game.get(team,{})
        pos,pos_source=resolve_position(b,avrow,pfrow,rolesrow,historical)
        model_role=str((pfrow or {}).get('model_role') or (pfrow or {}).get('role') or '')
        depth_role=str((avrow or {}).get('depth_chart_role') or (avrow or {}).get('raw_depth_role') or (rolesrow or {}).get('depth_chart_role') or '')
        avail='UNAVAILABLE' if avrow and bol(avrow.get('definitive_unavailable')) else str((avrow or {}).get('final_availability_state') or ('UNMATCHED_CURRENT_ROSTER' if not avrow else 'AVAILABLE'))
        match=avmatch if avmatch!='UNMATCHED' else (pfmatch if pfmatch!='UNMATCHED' else rolesmatch)
        market=str(b.get('source_market',''))
        li=lin.get((market,market_family(pos,market)),lin.get((market,'ALL'),{}))
        science=str(li.get('scientific_status','LINEAGE_NOT_RESOLVED'))
        po=num(O.get('fair_prob')); pu=num(U.get('fair_prob'))
        oo=num(b.get('vegas_over_odds')); uo=num(b.get('vegas_under_odds'))
        eo=ev_roi(po,oo); eu=ev_roi(pu,uo)
        side='' if eo is None and eu is None else ('OVER' if eu is None or (eo is not None and eo>=eu) else 'UNDER')
        ro,ru=implied_prob(oo),implied_prob(uo); nvo,nvu=no_vig(ro,ru),no_vig(ru,ro)
        bmp=po if side=='OVER' else pu if side=='UNDER' else None
        bmk=nvo if side=='OVER' else nvu if side=='UNDER' else None
        rows.append({
            'event_id':str(b.get('event_id','')),'game_id':str(gm.get('game_id','')),'kickoff':str(gm.get('kickoff_utc','')),
            'player':str(b.get('player','')),'key':key,'team':team,'opp':str(b.get('opponent','')),
            'pos':pos,'position_source':pos_source,'model_role':model_role,'depth_role':depth_role,
            'availability':avail,'authority':str((avrow or {}).get('availability_authority','')),'match':match,
            'game_status':str(gm.get('certification_state','UNMATCHED_GAME')),'game_eligible':bol(gm.get('production_eligible')),
            'market':market,'market_label':MLABEL.get(market,market),'book':str(b.get('book_title') or b.get('book','')),
            'line':num(b.get('vegas_line')),'proj':num(b.get('model_proj')),'sd':num(b.get('model_sd')),
            'over_odds':oo,'under_odds':uo,'p_over':po,'p_under':pu,'nv_over':nvo,'nv_under':nvu,
            'ev_over':eo,'ev_under':eu,'best_side':side,'best_ev':eo if side=='OVER' else eu if side=='UNDER' else None,
            'best_model_p':bmp,'best_market_p':bmk,'prob_edge':None if bmp is None or bmk is None else bmp-bmk,
            'science':science,'mean_owner':str(li.get('final_mean_owner','')),'limitation':str(li.get('known_limitation','')),
        })
    return rows


def signal(row,current):
    if row['pos']=='UNRESOLVED':
        return 'BLOCKED','BLOCKED'
    if 'NOT_DEDICATED' in row['science'] or row['science']=='LINEAGE_NOT_RESOLVED':
        return 'RESEARCH ONLY','BLOCKED'
    if row['game_status']=='KICKED_OFF_LOCKED' or row['availability'] in {'UNAVAILABLE','UNMATCHED_CURRENT_ROSTER'}:
        return 'BLOCKED','BLOCKED'
    e=row['best_ev']; q=row['prob_edge']
    sig='STRONG EDGE' if e is not None and q is not None and e>=.05 and q>=.03 else ('LEAN EDGE' if e is not None and e>0 else 'NO EDGE')
    if not current:
        return sig,'NO LIVE ODDS'
    return sig,('PLAY '+row['best_side'] if sig=='STRONG EDGE' else ('LEAN '+row['best_side'] if sig=='LEAN EDGE' else 'PASS'))


def parse_args():
    p=argparse.ArgumentParser()
    p.add_argument('--root',default='.')
    p.add_argument('--out',default='outputs/NFL_BETTING_MODEL_MASTER.xlsx')
    p.add_argument('--run-id',default=os.getenv('GITHUB_RUN_ID',''))
    p.add_argument('--sha',default=os.getenv('GITHUB_SHA',''))
    p.add_argument('--ref-name',default=os.getenv('GITHUB_REF_NAME',''))
    p.add_argument('--fetch-live-odds',default=os.getenv('FETCH_LIVE_ODDS','false'))
    return p.parse_args()


def main():
    a=parse_args(); root=Path(a.root); out=root/a.out; out.parent.mkdir(parents=True,exist_ok=True)
    pr=csvp(root/'outputs/props_priced_clean.csv')
    av=csvp(root/'data/current_player_availability.csv')
    cert=csvp(root/'data/current_player_availability_game_certification.csv')
    pf=csvp(root/'data/player_form.csv')
    roles=csvp(root/'data/roles_ourlads.csv')
    totals=csvp(root/'data/player_season_totals.csv')
    logs=csvp(root/'data/player_game_logs.csv')
    lineage=csvp(root/'data/market_model_lineage_current.csv')
    live=jsonp(root/'data/live_odds_status.json')
    requested=bol(a.fetch_live_odds)
    status='CURRENT' if requested and bol(live.get('available')) and not pr.empty else ('NO_LIVE_ODDS_REQUESTED' if not requested else 'NO_ACTIVE_MARKETS')
    rows=build_offer_rows(pr,av,pf,roles,cert,lineage,totals,logs)
    current=status=='CURRENT'
    unresolved=sum(1 for r in rows if r['pos']=='UNRESOLVED')

    wb=Workbook(); wb.calculation.fullCalcOnLoad=True; wb.calculation.forceFullCalc=True; wb.calculation.calcMode='auto'
    dash=wb.active; dash.title='Dashboard'; dash.merge_cells('A1:H2'); dash['A1']='NFL BETTING MODEL MASTER'; dash['A1'].fill=PatternFill('solid',fgColor=NAVY); dash['A1'].font=Font(bold=True,color=WHITE,size=16)
    dash.merge_cells('A4:H5')
    dash['A4']='CURRENT sportsbook prices are attached to this same Full Slate run.' if current else ('Live odds were not requested; football outputs are current and betting decisions are suppressed.' if not requested else 'Live odds were requested but no active pricing was available; betting decisions are suppressed.')
    dash['A4'].fill=PatternFill('solid',fgColor=GREEN if current else YELLOW); dash['A4'].font=Font(bold=True,color=GD if current else YD); dash['A4'].alignment=Alignment(wrap_text=True)
    kpis=[('Sportsbook Offers',len(rows)),('Unique Player-Markets',len({(r['key'],r['market']) for r in rows})),('Priced Players',len({r['key'] for r in rows if r['key']})),('Eligible Games',sum(bol(x) for x in cert.get('production_eligible',pd.Series(dtype=object)))),('Locked Games',sum(str(x)=='KICKED_OFF_LOCKED' for x in cert.get('certification_state',pd.Series(dtype=object)))),('Unresolved Position Rows',unresolved)]
    for i,(label,value) in enumerate(kpis):
        rr=7+(i//3)*3; cc=1+(i%3)*2; dash.cell(rr,cc,label).fill=PatternFill('solid',fgColor=LBLUE); dash.cell(rr,cc).font=Font(bold=True,color=NAVY); dash.cell(rr+1,cc,value).font=Font(bold=True,color=NAVY,size=16)
    dash['A14']='Pricing Status'; dash['B14']=status; dash.merge_cells('A18:H21'); dash['A18']='Sportsbook is downstream only. Position is resolved from football identity sources before any betting decision. Unresolved player identity/position is fail-closed at the row level.'; dash['A18'].fill=PatternFill('solid',fgColor=GRAY); dash['A18'].font=Font(color=GRAYD); dash['A18'].alignment=Alignment(wrap_text=True)
    for c in range(1,9): dash.column_dimensions[get_column_letter(c)].width=20

    master=wb.create_sheet('Master Betting Board')
    headers=['Event ID','Game ID','Kickoff UTC','Player','Player Key','Team','Opponent','Position','Position Source','Model Role','Depth Role','Current Availability','Availability Authority','Roster Match','Game Status','Game Eligible','Market','Book','Vegas Line','Model Projection','Model SD','Projection - Line','Over Odds','Under Odds','Model P Over','Model P Under','No-Vig P Over','No-Vig P Under','EV ROI Over','EV ROI Under','Best Side','Best Odds','Best Model P','Best Market P','Probability Edge','Best EV ROI','Snapshot Signal','Science Status','Bettable Now','Decision','Mean Owner / Limitation']
    master.append(headers)
    for r in rows:
        sg,dec=signal(r,current)
        master.append([r['event_id'],r['game_id'],r['kickoff'],r['player'],r['key'],r['team'],r['opp'],r['pos'],r['position_source'],r['model_role'],r['depth_role'],r['availability'],r['authority'],r['match'],r['game_status'],r['game_eligible'],r['market_label'],r['book'],r['line'],r['proj'],r['sd'],None,r['over_odds'],r['under_odds'],r['p_over'],r['p_under'],r['nv_over'],r['nv_under'],r['ev_over'],r['ev_under'],r['best_side'],r['over_odds'] if r['best_side']=='OVER' else r['under_odds'] if r['best_side']=='UNDER' else '',r['best_model_p'],r['best_market_p'],r['prob_edge'],r['best_ev'],sg,r['science'],current and r['game_eligible'] and r['pos']!='UNRESOLVED' and sg not in {'BLOCKED','RESEARCH ONLY'},dec,(r['mean_owner']+(' | '+r['limitation'] if r['limitation'] else ''))[:32000]])
        rr=master.max_row; master.cell(rr,22,f'=IF(OR(T{rr}="",S{rr}=""),"",T{rr}-S{rr})')
    header(master); master.freeze_panes='D2'; widths(master,28); master.column_dimensions['AO'].width=70; add_table(master,'MasterBettingBoard')
    for col in ['Y','Z','AA','AB','AC','AD','AG','AH','AI','AJ']:
        for cell in master[col][1:]: cell.number_format='0.0%'
    if master.max_row>1:
        master.conditional_formatting.add(f'AK2:AK{master.max_row}',FormulaRule(formula=['$AK2="STRONG EDGE"'],fill=PatternFill('solid',fgColor=GREEN),font=Font(color=GD,bold=True)))
        master.conditional_formatting.add(f'AK2:AK{master.max_row}',FormulaRule(formula=['$AK2="LEAN EDGE"'],fill=PatternFill('solid',fgColor=YELLOW),font=Font(color=YD,bold=True)))
        master.conditional_formatting.add(f'AK2:AK{master.max_row}',FormulaRule(formula=['OR($AK2="BLOCKED",$AK2="RESEARCH ONLY")'],fill=PatternFill('solid',fgColor=RED),font=Font(color=RD,bold=True)))
        master.conditional_formatting.add(f'AJ2:AJ{master.max_row}',ColorScaleRule(start_type='min',start_color='FECACA',mid_type='percentile',mid_value=50,mid_color='FEF3C7',end_type='max',end_color='BBF7D0'))

    best=wb.create_sheet('Best Snapshot Edges')
    bh=['Player','Team','Opp','Pos','Position Source','Model Role','Market','Best Book','Vegas Line','Model Projection','Projection-Line','Best Side','Best Odds','Model P','Market P (No-Vig)','Probability Edge','Best EV ROI','Snapshot Signal','Current Availability','Game Status','Science Status','Bettable Now','Decision']
    best.append(bh); picks={}
    for r in rows:
        k=(r['key'],r['market']); score=r['best_ev'] if r['best_ev'] is not None else -999
        if k not in picks or score>(picks[k]['best_ev'] if picks[k]['best_ev'] is not None else -999): picks[k]=r
    def rank(r):
        return (r['pos']=='UNRESOLVED',('NOT_DEDICATED' in r['science'] or r['science']=='LINEAGE_NOT_RESOLVED'),r['game_status']=='KICKED_OFF_LOCKED' or r['availability'] in {'UNAVAILABLE','UNMATCHED_CURRENT_ROSTER'},-(r['best_ev'] if r['best_ev'] is not None else -999))
    for r in sorted(picks.values(),key=rank):
        sg,dec=signal(r,current); odds=r['over_odds'] if r['best_side']=='OVER' else r['under_odds'] if r['best_side']=='UNDER' else ''
        best.append([r['player'],r['team'],r['opp'],r['pos'],r['position_source'],r['model_role'],r['market_label'],r['book'],r['line'],r['proj'],None,r['best_side'],odds,r['best_model_p'],r['best_market_p'],r['prob_edge'],r['best_ev'],sg,r['availability'],r['game_status'],r['science'],current and r['game_eligible'] and r['pos']!='UNRESOLVED' and sg not in {'BLOCKED','RESEARCH ONLY'},dec])
        rr=best.max_row; best.cell(rr,11,f'=IF(OR(J{rr}="",I{rr}=""),"",J{rr}-I{rr})')
    header(best); best.freeze_panes='B2'; widths(best,28); add_table(best,'BestSnapshotEdges')
    for col in ['N','O','P','Q']:
        for cell in best[col][1:]: cell.number_format='0.0%'
    if best.max_row>1:
        best.conditional_formatting.add(f'R2:R{best.max_row}',FormulaRule(formula=['$R2="STRONG EDGE"'],fill=PatternFill('solid',fgColor=GREEN),font=Font(color=GD,bold=True)))
        best.conditional_formatting.add(f'R2:R{best.max_row}',FormulaRule(formula=['$R2="LEAN EDGE"'],fill=PatternFill('solid',fgColor=YELLOW),font=Font(color=YD,bold=True)))
        best.conditional_formatting.add(f'R2:R{best.max_row}',FormulaRule(formula=['OR($R2="BLOCKED",$R2="RESEARCH ONLY")'],fill=PatternFill('solid',fgColor=RED),font=Font(color=RD,bold=True)))

    frame(wb,'Game Certification',cert,'GameCertification')
    frame(wb,'Availability & Roles',av,'AvailabilityRoles')
    frame(wb,'Market Science',lineage,'MarketScience')
    notes=wb.create_sheet('Lineage & Notes'); notes.append(['NFL BETTING MODEL MASTER — RUN LINEAGE','']); notes.merge_cells('A1:B1'); notes['A1'].fill=PatternFill('solid',fgColor=NAVY); notes['A1'].font=Font(bold=True,color=WHITE,size=15)
    for k,v in [('Workbook builder','scripts/build_master_betting_workbook_v1.py -> master_betting_workbook_core_v2.py'),('GitHub Run ID',a.run_id),('GitHub SHA',a.sha),('GitHub Ref',a.ref_name),('Live odds requested',requested),('Pricing status',status),('Live odds status',live.get('status','')),('Unresolved position rows',unresolved),('Position resolution','Exact current identity -> suffix-insensitive current identity -> historical player position -> deterministic model/market inference -> UNRESOLVED fail-closed.'),('Architecture','Sportsbook is downstream only; workbook generation cannot alter football projections.'),('Decision policy','PLAY/LEAN is a display gate only; it is not a calibrated staking policy.')]:
        notes.append([k,v]); notes.cell(notes.max_row,1).fill=PatternFill('solid',fgColor=LBLUE); notes.cell(notes.max_row,1).font=Font(bold=True,color=NAVY)
    notes.column_dimensions['A'].width=34; notes.column_dimensions['B'].width=105

    wb._sheets=[wb[n] for n in ['Dashboard','Best Snapshot Edges','Master Betting Board','Game Certification','Availability & Roles','Market Science','Lineage & Notes']]
    wb.save(out)
    audit={'disposition':'MASTER_BETTING_WORKBOOK_PUBLISHED','output':str(out),'pricing_status':status,'priced_offer_rows':len(rows),'player_market_rows':len({(r['key'],r['market']) for r in rows}),'availability_rows':len(av),'game_certification_rows':len(cert),'unresolved_position_rows':unresolved,'position_resolution_version':'suffix_historical_model_fallback_v2','sportsbook_downstream_only':True}
    audit_path=root/'data/master_betting_workbook_audit.json'; audit_path.parent.mkdir(parents=True,exist_ok=True); json.dump(audit,open(audit_path,'w',encoding='utf-8'),indent=2,sort_keys=True)
    print(json.dumps(audit,sort_keys=True))


if __name__=='__main__':
    main()
